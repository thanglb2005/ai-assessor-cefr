import io
import json
from copy import copy
from pathlib import Path
import subprocess

from openpyxl import load_workbook

repo = Path('/home/thanglvc/Documents/HCMUTE/TLCN/ai-assessor-cefr')
target = repo / 'docs/evidence/tc2-3-ai-usage/AI Prompt Log.xlsx'
baseline = load_workbook(io.BytesIO(subprocess.check_output([
    'git', '-C', str(repo), 'show',
    '2b428ce:docs/evidence/tc2-3-ai-usage/AI Prompt Log.xlsx',
])))
current = load_workbook(target)
assert baseline.sheetnames == current.sheetnames
preserved = []
for name in baseline.sheetnames:
    before, after = baseline[name], current[name]
    end = 17 if name == 'Thắng' else before.max_row
    for row in before.iter_rows(max_row=end):
        for cell in row:
            actual = after[cell.coordinate]
            assert actual.value == cell.value, (name, cell.coordinate, 'value')
            for attribute in ('font', 'fill', 'border', 'alignment', 'protection', 'number_format'):
                assert copy(getattr(actual, attribute)) == copy(getattr(cell, attribute)), (
                    name, cell.coordinate, attribute)
            assert bool(actual.hyperlink) == bool(cell.hyperlink)
            if cell.hyperlink:
                assert actual.hyperlink.target == cell.hyperlink.target
    assert list(before.merged_cells.ranges) == list(after.merged_cells.ranges)
    preserved.append({'sheet': name, 'rows': f'1–{end}', 'values_styles_links_preserved': True})
records = []
sheet = current['Thắng']
for row in range(18, 1001):
    values = [sheet.cell(row, col).value for col in range(1, 7)]
    if not any(v is not None for v in values):
        continue
    assert values[0].strftime('%Y-%m-%d') == '2026-10-02'
    assert values[1:3] == ['Thắng', 'Codex']
    assert sheet.cell(row, 1).number_format == 'dd/mm/yyyy'
    paths = values[5].split('; ')
    assert all((repo / path).is_file() for path in paths)
    is_full_file = len(paths) == 1 and '/prompts/' in paths[0]
    if is_full_file:
        assert (repo / paths[0]).read_text() == values[4]
    else:
        assert row == 18
        assert values[4] == 'bạn hãy promt r gửi cho các subagent luna 6 đi làm task  hết của w3 đi, promt gì thì lưu vào các file theo chuẩn quy trình AI log, bạn là người reviewcode, tạo nhánh rieegn nha'
    records.append({'row': row, 'date': '2026-10-02', 'person': values[1], 'tool': values[2],
                    'activity': values[3], 'evidence': paths,
                    'full_text_matches_file': True if is_full_file else 'actual user instruction'})
assert len(records) == 12
assert sum(r['full_text_matches_file'] is True for r in records) == 11
summary_formulas = [cell.value for row in current['Tổng hợp'].iter_rows() for cell in row
                    if cell.data_type == 'f']
result = {
    'result': 'PASS', 'baseline_revision': '2b428ceea9e4459f3235d4ce67f10bab31585ebf',
    'preserved': preserved, 'new_records': records, 'summary_formulas_preserved': True,
    'summary_formula_count': len(summary_formulas), 'workbook_zip_parse': 'PASS',
    'limits': 'Kiểm workbook/cell/full text/path; không tự tạo Jira key/giờ công hoặc user verdict.',
}
output = repo / 'docs/sdd/features/w3-local-integration/evidence/final/ai-log-audit.json'
output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'result': 'PASS', 'new_records': len(records),
                  'full_file_text_matches': 11, 'preserved_sheets': len(preserved)}, ensure_ascii=False))
